"""
PDF/DOCX Parser Service — Extracts text from uploaded resume files.
Uses PyMuPDF for PDFs and python-docx for Word documents.
"""

import re
import fitz  # PyMuPDF
from docx import Document


def extract_text_from_pdf(file_path):
    """
    Extract all text from a PDF file using PyMuPDF.
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        str: Extracted and cleaned text
    """
    try:
        doc = fitz.open(file_path)
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return clean_text(text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_path):
    """
    Extract all text from a DOCX file using python-docx.
    
    Args:
        file_path: Path to the DOCX file
        
    Returns:
        str: Extracted and cleaned text
    """
    try:
        doc = Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        
        return clean_text(text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text(file_path, file_type='pdf'):
    """
    Extract text from a file based on its type.
    
    Args:
        file_path: Path to the file
        file_type: 'pdf' or 'docx'
        
    Returns:
        str: Extracted and cleaned text
    """
    if file_type == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_type == 'docx':
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def clean_text(text):
    """
    Clean and normalize extracted text.
    
    - Remove excessive whitespace
    - Remove special/control characters
    - Normalize line breaks
    """
    if not text:
        return ""
    
    # Remove null bytes and control characters (except newlines/tabs)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    
    # Replace multiple spaces with single space
    text = re.sub(r'[ \t]+', ' ', text)
    
    # Replace more than 2 consecutive newlines with 2
    text = re.sub(r'\n{3,}', '\n\n', text)
    
    # Strip leading/trailing whitespace from each line
    lines = [line.strip() for line in text.split('\n')]
    text = '\n'.join(lines)
    
    # Remove leading/trailing whitespace
    text = text.strip()
    
    return text
